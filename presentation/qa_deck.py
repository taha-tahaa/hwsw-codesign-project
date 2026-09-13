#!/usr/bin/env python3
"""Geometry and text-fit QA for the deck, without a renderer.

There is no LibreOffice or pdftoppm on either machine we have, so slides cannot
be rasterized and eyeballed.  This checks the defects that a visual pass would
normally catch, arithmetically:

  1. shapes outside the slide, or inside the 0.5" margin
  2. text that will not fit its box at the given font size
  3. text boxes that overlap each other

The text-fit model is deliberately pessimistic: it assumes an average glyph
width of 0.50 em (Calibri/Cambria run nearer 0.46-0.48 for mixed-case English),
so it flags borderline boxes rather than missing them.  Treat a WARN as "go
look", not as "definitely broken".

    python qa_deck.py deck.pptx
"""

import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.5
GLYPH_W = 0.50      # em per average character
LINE_H = 1.22       # line box as a multiple of font size


def inches(v):
    return None if v is None else v / EMU_IN


def fits(text, width_in, height_in, pt):
    """Estimate whether `text` fits a box, honouring explicit newlines."""
    if not text.strip():
        return True, 0, 0
    char_w = (pt * GLYPH_W) / 72.0
    line_h = (pt * LINE_H) / 72.0
    if char_w <= 0 or width_in <= 0:
        return True, 0, 0
    per_line = max(1, int(width_in / char_w))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, -(-len(para) // per_line))  # ceil
    return (lines * line_h) <= height_in + 1e-6, lines, lines * line_h


def max_pt(shape):
    pts = []
    if not shape.has_text_frame:
        return 12.0
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                pts.append(run.font.size.pt)
    return max(pts) if pts else 12.0


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ox = min(ax + aw, bx + bw) - max(ax, bx)
    oy = min(ay + ah, by + bh) - max(ay, by)
    if ox <= 0.02 or oy <= 0.02:
        return 0.0
    return ox * oy


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
    prs = Presentation(path)
    problems = 0
    warns = 0

    for idx, slide in enumerate(prs.slides, start=1):
        issues = []
        text_boxes = []

        for shape in slide.shapes:
            x, y = inches(shape.left), inches(shape.top)
            w, h = inches(shape.width), inches(shape.height)
            if None in (x, y, w, h):
                continue
            name = (shape.shape_type.__str__() if shape.shape_type else "shape")

            # 1. bounds
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                issues.append(("ERROR", "off-slide: %s at (%.2f,%.2f) %.2fx%.2f"
                               % (name, x, y, w, h)))
            elif x < MARGIN - 0.01 or y < MARGIN - 0.01 or \
                    x + w > SLIDE_W - MARGIN + 0.01 or y + h > SLIDE_H - MARGIN + 0.35:
                # bottom chrome (speaker/page number) legitimately sits low
                if y < 6.85:
                    issues.append(("WARN", "inside 0.5in margin: %s at (%.2f,%.2f) %.2fx%.2f"
                                   % (name, x, y, w, h)))

            # 2. text fit
            if shape.has_text_frame and shape.text_frame.text.strip():
                pt = max_pt(shape)
                ok, lines, need = fits(shape.text_frame.text, w, h, pt)
                if not ok:
                    snippet = shape.text_frame.text.strip().replace("\n", " ")[:52]
                    issues.append(("WARN",
                                   "may overflow: %.0fpt needs ~%.2fin, box %.2fin | \"%s\""
                                   % (pt, need, h, snippet)))
                text_boxes.append(((x, y, w, h), shape.text_frame.text.strip()[:28]))

        # 3. text-on-text overlap
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                a, ta = text_boxes[i]
                b, tb = text_boxes[j]
                ov = overlap(a, b)
                if ov > 0.12:
                    issues.append(("WARN", "text overlap %.2f sq-in: \"%s\" / \"%s\""
                                   % (ov, ta, tb)))

        if issues:
            print("slide %d" % idx)
            for lvl, msg in issues:
                print("   %-5s %s" % (lvl, msg))
                if lvl == "ERROR":
                    problems += 1
                else:
                    warns += 1

    print()
    print("slides: %d   errors: %d   warnings: %d"
          % (len(prs.slides), problems, warns))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
